"""End to end: the whole pipeline runs on the fixture and writes every output."""
from __future__ import annotations

import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import FIXTURES
from run_pipeline import run
from src.config import Paths
from src.star_schema import STAR_TABLES


@pytest.fixture(scope="module")
def outputs(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    paths = Paths(raw=FIXTURES, processed=out / "processed", reports=out / "reports", docs=out)
    result = run(paths, download=False)
    return paths, result


def test_star_schema_files(outputs):
    paths, _ = outputs
    for t in STAR_TABLES:
        csv, parquet = paths.processed / f"{t}.csv", paths.processed / f"{t}.parquet"
        assert csv.exists() and parquet.exists(), t
        assert len(pd.read_parquet(parquet)) == len(pd.read_csv(csv))


def test_star_schema_keys_resolve(outputs):
    """Every fact key must find its dimension row (no silent blanks in Power BI)."""
    _, result = outputs
    con = result["con"]
    for fact in ["fact_orders", "fact_reconciliation"]:
        for key, dim, dim_key in [("state_code", "dim_customer_state", "state_code"),
                                  ("payment_type", "dim_payment_type", "payment_type"),
                                  ("category_key", "dim_product_category", "category_key")]:
            missing = con.execute(f"select count(*) from {fact} f anti join {dim} d on f.{key} = d.{dim_key}").fetchone()[0]
            assert missing == 0, f"{fact}.{key}"
        missing = con.execute(f"""select count(*) from {fact} f anti join dim_date d on f.purchase_date_key = d.date_key
                                  where f.purchase_date_key is not null""").fetchone()[0]
        assert missing == 0


def test_fact_row_counts(outputs):
    _, result = outputs
    con = result["con"]
    assert con.execute("select count(*) from fact_orders").fetchone()[0] == \
        con.execute("select count(distinct order_id) from stg_orders").fetchone()[0]
    assert con.execute("select count(*) from fact_exceptions").fetchone()[0] == \
        con.execute("select count(*) from dq_exceptions").fetchone()[0] + \
        con.execute("select count(*) from recon_exceptions").fetchone()[0]


def test_metrics_json(outputs):
    paths, _ = outputs
    m = json.loads((paths.reports / "metrics.json").read_text())
    rec = m["reconciliation"]
    assert rec["orders_matched"] + rec["orders_unmatched"] == rec["orders_total"]
    assert sum(c["orders"] for c in rec["categories"]) == rec["orders_unmatched"]
    assert m["exceptions"]["total"] == m["data_quality"]["exceptions"] + rec["orders_unmatched"]
    assert m["currency"] == "BRL"


def test_reports_exist(outputs):
    paths, _ = outputs
    wb = load_workbook(paths.reports / "summary.xlsx")
    assert {"KPIs", "Exceptions", "Reconciliation"} <= set(wb.sheetnames)
    page = (paths.reports / "dashboard.html").read_text()
    assert "<title>Olist Reconciliation Dashboard</title>" in page and "Plotly" in page
