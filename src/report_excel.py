"""reports/summary.xlsx: a workbook a manager can open without any tools.

Sheets
  KPIs            headline numbers with a plain-language definition each
  Gap by category reconciliation buckets: orders, BRL, rule, evidence, action
  Exceptions      summary by check and severity, then every exception row
  Reconciliation  every unmatched order, biggest gap first, with ranks
  Monthly         monthly trend table
  States / Categories  where mismatches concentrate
  Forecast        holdout scores and the next 3 months
"""
from __future__ import annotations

import decimal
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRL = '"R$" #,##0.00'
INT = "#,##0"
PCT = '0.00"%"'
HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="52514E")
THIN = Border(bottom=Side(style="thin", color="E1E0D9"))
SEVERITY_FILLS = {"high": "F8D7D7", "medium": "FDEBD3", "low": "FFF6D6"}


def _fmt_for(col: str) -> str | None:
    c = col.lower()
    if c.endswith("_brl") or c in {"expected_total", "paid_total", "gap", "abs_gap", "product_value", "freight_value"}:
        return BRL
    if c.endswith("_pct"):
        return PCT
    if c in {"orders", "unmatched", "matched", "exceptions", "delivered", "actual", "predicted", "forecast_orders",
             "distinct_records", "rows_in", "clean", "warning", "quarantined"}:
        return INT
    return None


def _write_table(ws, df: pd.DataFrame, start_row: int, title: str | None = None, note: str | None = None) -> int:
    r = start_row
    if title:
        ws.cell(r, 1, title).font = TITLE_FONT
        r += 1
    if note:
        ws.cell(r, 1, note).font = NOTE_FONT
        r += 1
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(r, j, col.replace("_", " "))
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    header_row = r
    for row in df.itertuples(index=False):
        r += 1
        for j, v in enumerate(row, 1):
            if v is pd.NA or v is pd.NaT or (isinstance(v, float) and pd.isna(v)):
                v = None
            elif isinstance(v, decimal.Decimal):
                v = float(v)
            elif hasattr(v, "item"):  # numpy scalar
                v = v.item()
            cell = ws.cell(r, j, v)
            fmt = _fmt_for(df.columns[j - 1])
            if fmt:
                cell.number_format = fmt
            cell.border = THIN
    if "severity" in df.columns and len(df):
        col = get_column_letter(list(df.columns).index("severity") + 1)
        rng = f"{col}{header_row + 1}:{col}{r}"
        for sev, color in SEVERITY_FILLS.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{sev}"'],
                                                          fill=PatternFill("solid", fgColor=color)))
    return r + 2


def _autosize(ws, max_width: int = 60) -> None:
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                n = len(str(c.value)) if not isinstance(c.value, float) else 14
                widths[c.column_letter] = max(widths.get(c.column_letter, 8), min(n + 2, max_width))
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_excel(con: duckdb.DuckDBPyConnection, metrics: dict, kpis: dict[str, pd.DataFrame], path: Path) -> None:
    rec, dq, dlv = metrics["reconciliation"], metrics["data_quality"], metrics["kpis"]["delivery"]
    reviews = metrics["kpis"]["reviews_by_match"]
    fc = metrics["forecast"]
    wb = Workbook()

    # KPIs
    ws = wb.active
    ws.title = "KPIs"
    ws["A1"] = "Olist reconciliation, data quality and KPI summary"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"Orders purchased {metrics['dataset']['first_purchase_date']} to "
                f"{metrics['dataset']['last_purchase_date']}. Currency: Brazilian real (BRL, R$). "
                f"Generated {metrics['generated_at']}.")
    ws["A2"].font = NOTE_FONT
    rows = [
        ("Orders reconciled", rec["orders_total"], INT, "Every order_id found in orders, items or payments"),
        ("Match rate", rec["match_rate_pct"], PCT, f"Orders where paid = sold within R$ {metrics['match_tolerance_brl']:.2f}"),
        ("Match rate (orders in both systems)", rec["match_rate_both_systems_pct"], PCT,
         "Same, only orders that have both items and payments"),
        ("Unmatched orders", rec["orders_unmatched"], INT, "Orders outside tolerance or missing a side"),
        ("Total absolute gap", rec["abs_gap_brl"], BRL, "Sum of |paid - sold| over unmatched orders"),
        ("Average absolute gap per unmatched order", rec["avg_abs_gap_brl"], BRL, ""),
        ("Gap as % of paid value", rec["abs_gap_pct_of_paid"], '0.000"%"', "Total absolute gap / total paid"),
        ("Paid on orders with no items", rec["missing_items"]["paid_brl"], BRL,
         "Payments with no order items: refund exposure"),
        ("Unexplained orders", rec["orders_unexplained"], INT, "Fit no evidence-based bucket: manual review"),
        ("High-severity reconciliation orders", rec["orders_high_severity"], INT, ""),
        ("Data quality checks run", dq["checks_run"], INT, "Schema contract + business rules"),
        ("Data quality exceptions", dq["exceptions"], INT, "Rows are flagged, never deleted"),
        ("Rows with no data quality exception", dq["pct_rows_clean"], PCT, "Across all 8 raw tables"),
        ("On-time delivery rate", dlv["on_time_rate_pct"], PCT, "Delivered on or before the estimated date"),
        ("Late delivery rate", dlv["late_rate_pct"], PCT, ""),
        ("Avg review score, matched orders", reviews["matched"]["avg_review_score"], "0.00", "1 to 5 stars"),
        ("Avg review score, unmatched orders", reviews["unmatched"]["avg_review_score"], "0.00", "1 to 5 stars"),
    ]
    if not fc.get("skipped"):
        rows += [
            ("Forecast model chosen", fc["winner"], None, "Lowest holdout MAPE"),
            ("Holdout MAPE of chosen model", fc["winner_mape_pct"], PCT, f"Holdout months {', '.join(fc['holdout_months'])}"),
        ]
    for j, h in enumerate(["KPI", "Value", "Definition"], 1):
        c = ws.cell(4, j, h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for i, (name, value, fmt, note) in enumerate(rows, 5):
        ws.cell(i, 1, name).font = Font(bold=True)
        c = ws.cell(i, 2, value)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        ws.cell(i, 3, note).font = NOTE_FONT
        for j in range(1, 4):
            ws.cell(i, j).border = THIN
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 70
    ws.freeze_panes = "A5"

    # Gap by category
    ws = wb.create_sheet("Gap by category")
    summary = con.execute("select * from reconciliation_summary").df()
    _write_table(ws, summary, 1, "Reconciliation buckets",
                 "Every order is in exactly one bucket. Rules run top to bottom; evidence files are in reports/evidence/.")
    _autosize(ws)
    ws.freeze_panes = "B4"

    # Exceptions
    ws = wb.create_sheet("Exceptions")
    r = _write_table(ws, kpis["exceptions_by_check"], 1, "Exceptions by check",
                     "Data quality checks + reconciliation. Severity: high = money or counts wrong, "
                     "medium = suspicious, low = informational.")
    detail = con.execute("""
        select exception_id, source, table_name, record_key, check_name, severity, reason
        from fact_exceptions
        order by case severity when 'high' then 1 when 'medium' then 2 else 3 end, source, check_name, record_key
    """).df()
    _write_table(ws, detail, r, "Every exception", None)
    _autosize(ws)

    # Reconciliation
    ws = wb.create_sheet("Reconciliation")
    unmatched = con.execute("""
        select k.overall_rank, r.order_id, r.order_status, r.purchase_date_key, r.state_code, r.payment_types,
               r.max_installments, r.expected_total, r.paid_total, r.gap, r.abs_gap, r.gap_pct, r.gap_category,
               r.severity, r.gap_rank_in_month, r.gap_rank_in_state,
               round(100 * r.cumulative_gap_share, 2) as cumulative_gap_share_pct
        from fact_reconciliation r
        join reconciliation_ranked k using (order_id)
        where not r.is_matched
        order by k.overall_rank
    """).df()
    _write_table(ws, unmatched, 1, "Unmatched orders, largest gap first",
                 "expected = items price + freight (system A); paid = payments (system B); gap = paid - expected.")
    _autosize(ws, 34)
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(unmatched.shape[1])}{3 + len(unmatched)}"

    # Monthly, states, categories
    for name, key, title in [("Monthly", "monthly", "Monthly trend"),
                             ("States", "top_states", "Customer states by unmatched orders"),
                             ("Categories", "top_categories", "Product categories by unmatched orders (orders with items)")]:
        ws = wb.create_sheet(name)
        _write_table(ws, kpis[key], 1, title)
        _autosize(ws)
        ws.freeze_panes = "B3"

    # Forecast
    ws = wb.create_sheet("Forecast")
    if fc.get("skipped"):
        ws["A1"] = f"Forecast skipped: {fc['reason']}"
    else:
        r = _write_table(ws, pd.DataFrame(fc["scores"]), 1, "Holdout accuracy (lower MAPE is better)",
                         f"Trained on {fc['window_start']} to {fc['holdout_months'][0]} (exclusive); "
                         f"excluded incomplete months: {', '.join(fc['excluded_months'])}.")
        r = _write_table(ws, pd.DataFrame(fc["holdout"]), r, "Holdout predictions vs actual")
        _write_table(ws, pd.DataFrame(fc["forecast"]), r, f"Next 3 months: {fc['winner']}",
                     "A baseline won." if fc["baseline_wins"] else "A model beat the baselines.")
    _autosize(ws)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
